#!/usr/bin/env python3
"""Rebuild the vector PDF, editable PowerPoint, previews, and layout audit.

python build_deck.py
python build_deck.py --screenshot /absolute/path/to/verified-results.png
"""
from pathlib import Path
import argparse
import json
import math

from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image, ImageDraw
import pymupdf as fitz

ROOT = Path(__file__).resolve().parent
W, H = 960, 540
BG = '#F7F5EF'
INK = '#172E30'
DARK = '#132D2F'
MUTED = '#536665'
RULE = '#CDD4CE'
WHITE = '#FFFFFF'
CORAL = '#F87759'
ORANGE = '#BD452B'
MINT = '#B9E5CE'
PALE = '#E6EDE5'
LIGHTMUTED = '#C3D0CB'
PANEL = '#203D3E'

for face in ['Regular', 'SemiBold', 'Bold']:
    pdfmetrics.registerFont(TTFont(face, str(ROOT / 'assets/fonts' / f'Pretendard-{face}.ttf')))

NOTES = [
    {
        'time': '00:00–00:20', 'seconds': 20, 'title': '느린 코드를 직접 고치는 AI',
        'ko': '안녕하세요. AgenticRocket은 느린 코드의 병목을 찾아 직접 고치는 AI 성능 최적화 솔루션입니다. 프로그램 안쪽의 로우레벨 코드를 다룹니다. 불필요한 계산이나 메모리 작업을 줄여 같은 일을 더 빠르게 끝내는 것이 목표입니다.',
        'en': 'AgenticRocket is an AI performance engineer that finds bottlenecks and changes the code. We focus on low-level optimization: removing unnecessary computation and memory work so software can do the same job faster.',
        'cue': '첫 문장 뒤 한 박자 쉰다. “직접 고치는”에 힘을 준다.'
    },
    {
        'time': '00:20–00:55', 'seconds': 35, 'title': 'Wine으로 이해하는 로우레벨 최적화',
        'ko': 'Wine을 예로 들어보겠습니다. Windows 프로그램을 Linux에서도 실행하게 해주는 오픈소스입니다. 프로그램은 잘 실행되는데, 안에서 같은 메모리 공간을 매번 새로 준비하느라 시간을 쓴다고 가정해보죠. 기존 공간을 다시 쓸 수 있다면, 같은 결과를 내는 데 필요한 작업이 줄어듭니다. 화면은 이런 최적화를 설명하는 예시입니다. AgenticRocket은 이렇게 코드 깊숙한 곳의 낭비를 찾아 줄이려 합니다.',
        'en': 'Take Wine, the open-source project that runs Windows programs on Linux. Imagine a frequently called path allocating fresh memory every time. If it can safely reuse that memory, it can do the same work with less overhead. This diagram illustrates an optimization idea, rather than a measured Wine result. It makes our focus tangible: finding and removing wasted work deep inside the code.',
        'cue': '왼쪽의 반복 할당에서 오른쪽의 재사용으로 손짓한다. Wine은 활용 시나리오로 소개한다.'
    },
    {
        'time': '00:55–01:30', 'seconds': 35, 'title': '병목 발견부터 실제 소스 수정까지',
        'ko': '사용자는 프로젝트와 목표를 맡깁니다. 기존 동작을 유지하면서 더 빠르게 만들어달라는 겁니다. 에이전트는 코드를 읽고 실행 결과를 확인해 느린 원인을 찾습니다. 수정할 방법을 정하면 소스를 직접 바꾸고 빌드와 테스트를 돌립니다. 오류가 나면 그 결과를 읽고 다시 고칩니다. 이 일을 한 에이전트가 같은 프로젝트에서 계속 맡습니다. 사용자가 이 패치로 다시 측정해달라고 하면, 앞선 작업을 이어서 진행합니다.',
        'en': 'You give the agent a project and a goal: make it faster while preserving its behavior. It reads the code, examines execution results, and narrows down the bottleneck. Then it edits the source, builds it, and runs the tests. If something fails, it uses that result to revise the fix. One agent stays with the project, so follow-up requests continue from the existing work.',
        'cue': '분석 → 수정 → 재실행 순서로 짚는다. “소스를 직접 바꾼다”가 핵심이다.'
    },
    {
        'time': '01:30–02:10', 'seconds': 40, 'title': 'Daytona에서 실행하고 반복 검증',
        'ko': '여기서 Daytona가 핵심 역할을 합니다. AI가 실제 코드를 고치고 실행하는 작업 공간이면서 개선을 확인하는 실험실입니다. 수정안 하나를 세 개의 샌드박스, 즉 분리된 실행 공간에 보냅니다. 각 공간 안에서 원본과 수정본을 같은 조건으로 여러 번 비교합니다. 기능 테스트와 결과값도 확인합니다. 한 번 우연히 빨랐다는 이유로 성공이라고 하지 않습니다. 반복해서 개선이 확인돼야 하고 결과가 불안정하면 판단을 보류합니다.',
        'en': 'Daytona provides both the workspace where the agent changes and runs code, and the sandboxes where we check the improvement. We send one fixed patch to three sandboxes. Within each sandbox, we compare the original and modified code repeatedly under matched conditions. We also check tests and outputs. One lucky fast run is not enough. The improvement must be consistent, and uncertain results remain inconclusive.',
        'cue': '세 상자를 천천히 짚는다. “각 공간 안에서 전후 비교”를 분명히 말한다.'
    },
    {
        'time': '02:10–02:40', 'seconds': 30, 'title': '코드와 검증 근거를 함께 전달',
        'ko': '개발자는 실제 수정 코드와 검증 근거를 가져갑니다. 무엇을 바꿨는지, 어떤 기능 테스트를 통과했는지, 측정한 작업이 얼마나 빨라졌는지 함께 확인합니다. 수정안과 보고서를 내려받아 코드 리뷰에 사용할 수 있습니다. AI가 왜 이렇게 고쳤는지도 같은 프로젝트에서 물어볼 수 있습니다. 개발자는 코드와 실행 결과를 보고 이 변경을 받아들일지 결정합니다.',
        'en': 'The developer gets the actual patch together with the evidence: what changed, which tests passed, and the measured performance difference. The patch and report can be downloaded for code review. You can also ask the same agent why it made the change. That gives the developer something concrete to review before deciding whether to adopt the optimization.',
        'cue': '실제 화면이 준비되면 오른쪽 보고서 도식을 교체하고, 수정 코드와 측정 결과를 짚는다.'
    },
    {
        'time': '02:40–03:00', 'seconds': 20, 'title': 'AgenticRocket',
        'ko': '저희가 만들고 싶은 것은, 프로젝트를 맡기면 병목을 찾아 고치고 검증까지 이어가는 AI 성능 엔지니어입니다. Wine 같은 복잡한 소프트웨어에도 이런 최적화 작업을 더 쉽게 맡길 수 있도록 하겠습니다. AgenticRocket. AI가 고치고 실행으로 증명합니다. 감사합니다.',
        'en': 'Our goal is an AI performance engineer that stays with a project, finds the bottleneck, makes the fix, and verifies it. We want to make that work accessible even for complex software like Wine. AgenticRocket: AI makes the fix. Real runs prove the gain. Thank you.',
        'cue': '마지막 제품명과 핵심 문장을 말한 뒤 멈춘다.'
    }
]


def rgb(value):
    return RGBColor.from_string(value.lstrip('#'))


class Deck:
    def __init__(self):
        self.pdf = canvas.Canvas(str(ROOT / 'AgenticRocket_3min_KO-EN.pdf'), pagesize=(W, H), pageCompression=1, initialFontName='Regular', initialFontSize=12)
        self.pdf.setTitle('AgenticRocket — AI low-level performance optimization | 한국어 · English')
        self.pdf.setAuthor('AgenticRocket')
        self.pdf.setSubject('Three-minute bilingual pitch. Wine optimization scenario and Daytona validation.')
        self.pdf.setViewerPreference('DisplayDocTitle', 'true')
        self.ppt = Presentation()
        self.ppt.slide_width = Pt(W)
        self.ppt.slide_height = Pt(H)
        self.ppt.core_properties.title = 'AgenticRocket — 3-minute bilingual pitch'
        self.ppt.core_properties.subject = 'AI low-level performance optimization'
        self.ppt.core_properties.author = 'AgenticRocket'
        self.audit = []
        self.index = 0
        self.slide = None

    def page(self, dark=False, section=''):
        if self.index:
            self.pdf.showPage()
        self.index += 1
        self.slide = self.ppt.slides.add_slide(self.ppt.slide_layouts[6])
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = rgb(DARK if dark else BG)
        self.pdf.setFillColor(HexColor(DARK if dark else BG))
        self.pdf.rect(0, 0, W, H, fill=1, stroke=0)
        self.fore = BG if dark else INK
        self.muted = LIGHTMUTED if dark else MUTED
        self.rule = '#3D5655' if dark else RULE
        self.accent = CORAL if dark else ORANGE
        self.line(56, 494, 904, 494, self.rule, 0.8)
        self.text('AGENTICROCKET', 56, 508, 10, self.muted, 'SemiBold', 260, essential=False)
        self.text(section, 352, 508, 10, self.muted, 'Regular', 400, essential=False)
        self.text(f'{self.index:02d} / 06', 852, 508, 10, self.muted, 'Regular', 65, essential=False)
        note = NOTES[self.index - 1]
        self.slide.notes_slide.notes_text_frame.text = (
            f"{note['time']} | {note['seconds']} seconds\n\n한국어\n{note['ko']}\n\nEnglish\n{note['en']}\n\n발표 팁\n{note['cue']}"
        )

    def text(self, value, x, y, size, color=None, face='Regular', width=None, leading=1.14, essential=True, name=None):
        # Requested refinement: quieter typography, with readable body text.
        if size >= 80:
            size = round(size * 0.88)
        elif size >= 40:
            size = round(size * 0.92)
        elif size >= 30:
            size = round(size * 0.94)
        elif size >= 25:
            size = max(24, round(size * 0.96))
        if essential:
            size = max(24, size)
        color = color or self.fore
        lines = value.split('\n')
        widths = [pdfmetrics.stringWidth(v, face, size) for v in lines]
        real_w = max(widths, default=0)
        box_w = width if width is not None else real_w + 3
        if real_w > box_w + 0.1:
            raise ValueError(f'Slide {self.index} text overflow: {value!r}: {real_w:.1f} > {box_w:.1f}')
        asc = pdfmetrics.getAscent(face) / 1000 * size
        desc = abs(pdfmetrics.getDescent(face) / 1000 * size)
        real_h = (len(lines) - 1) * size * leading + asc + desc
        if x < 0 or y < 0 or x + real_w > W + 0.1 or y + real_h > H + 0.1:
            raise ValueError(f'Slide {self.index} out-of-bounds text: {value!r}')
        self.pdf.setFillColor(HexColor(color))
        self.pdf.setFont(face, size)
        for i, val in enumerate(lines):
            self.pdf.drawString(x, H - y - asc - i * size * leading, val)
        shape = self.slide.shapes.add_textbox(Pt(x), Pt(y), Pt(box_w + 2), Pt(max(real_h + 8, size * 1.3)))
        if name:
            shape.name = name
        tf = shape.text_frame
        tf.clear()
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, val in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = p.space_after = Pt(0)
            p.line_spacing = Pt(size * leading)
            r = p.add_run()
            r.text = val
            r.font.name = 'Pretendard'
            r.font.size = Pt(size)
            r.font.bold = face in ['Bold', 'SemiBold']
            r.font.color.rgb = rgb(color)
            rpr = r._r.get_or_add_rPr()
            rpr.set('lang', 'ko-KR' if any('\uac00' <= c <= '\ud7a3' for c in val) else 'en-US')
            ea = OxmlElement('a:ea')
            ea.set('typeface', 'Pretendard')
            rpr.append(ea)
        self.audit.append({'page': self.index, 'text': value, 'bbox': [x, y, real_w, real_h], 'font_pt': size, 'essential': essential})

    def rect(self, x, y, w, h, fill=None, stroke=None, radius=0, line_width=1, name=None):
        self.pdf.setLineWidth(line_width)
        if fill:
            self.pdf.setFillColor(HexColor(fill))
        if stroke:
            self.pdf.setStrokeColor(HexColor(stroke))
        if radius:
            self.pdf.roundRect(x, H - y - h, w, h, radius, fill=bool(fill), stroke=bool(stroke))
        else:
            self.pdf.rect(x, H - y - h, w, h, fill=bool(fill), stroke=bool(stroke))
        sh = self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
        if name:
            sh.name = name
        if radius:
            sh.adjustments[0] = min(radius / min(w, h), 0.5)
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(fill)
        else:
            sh.fill.background()
        if stroke:
            sh.line.color.rgb = rgb(stroke)
            sh.line.width = Pt(line_width)
        else:
            sh.line.fill.background()

    def line(self, x1, y1, x2, y2, color, thickness=1):
        self.pdf.setStrokeColor(HexColor(color))
        self.pdf.setLineWidth(thickness)
        self.pdf.setLineCap(1)
        self.pdf.line(x1, H - y1, x2, H - y2)
        sh = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
        sh.line.color.rgb = rgb(color)
        sh.line.width = Pt(thickness)

    def arrow(self, x1, y1, x2, y2, color, thickness=2.5, tip=8):
        self.line(x1, y1, x2, y2, color, thickness)
        angle = math.atan2(y2-y1, x2-x1)
        for d in [-0.65, 0.65]:
            self.line(x2, y2, x2 - tip*math.cos(angle+d), y2 - tip*math.sin(angle+d), color, thickness)

    def circle(self, x, y, diameter, fill, stroke=None, line_width=1):
        self.pdf.setFillColor(HexColor(fill))
        if stroke:
            self.pdf.setStrokeColor(HexColor(stroke))
        self.pdf.setLineWidth(line_width)
        self.pdf.circle(x+diameter/2, H-y-diameter/2, diameter/2, fill=1, stroke=bool(stroke))
        sh = self.slide.shapes.add_shape(MSO_SHAPE.OVAL, Pt(x), Pt(y), Pt(diameter), Pt(diameter))
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
        if stroke:
            sh.line.color.rgb = rgb(stroke)
            sh.line.width = Pt(line_width)
        else:
            sh.line.fill.background()

    def mark(self, x, y, size=24, color=None):
        color = color or self.accent
        self.line(x, y+size, x+size, y, color, size*0.15)
        self.line(x+size*0.5, y, x+size, y, color, size*0.15)
        self.line(x+size, y, x+size, y+size*0.5, color, size*0.15)

    def label(self, ko, en, x=56, y=40):
        self.text(ko, x, y, 17, self.muted, 'SemiBold', essential=False)
        offset = pdfmetrics.stringWidth(ko, 'SemiBold', 17) + 18
        self.text(en, x+offset, y+1, 16, self.muted, 'Regular', essential=False)

    def image(self, path, x, y, w, h):
        im = Image.open(path)
        scale = min(w/im.width, h/im.height)
        rw, rh = im.width*scale, im.height*scale
        xx, yy = x+(w-rw)/2, y+(h-rh)/2
        self.pdf.drawImage(str(path), xx, H-yy-rh, width=rw, height=rh, mask='auto')
        self.slide.shapes.add_picture(str(path), Pt(xx), Pt(yy), width=Pt(rw), height=Pt(rh))

    def save(self):
        self.pdf.save()
        self.ppt.save(str(ROOT/'AgenticRocket_3min_KO-EN.pptx'))
        (ROOT/'layout-audit.json').write_text(json.dumps(self.audit, ensure_ascii=False, indent=2))


def build(screenshot=None):
    d = Deck()

    # 1 — The product is an optimizer, introduced in the first sentence.
    d.page(True, 'AI PERFORMANCE ENGINEER')
    d.mark(57, 39, 20)
    d.text('AgenticRocket', 91, 33, 27, BG, 'SemiBold')
    d.text('느린 코드,', 56, 119, 61, BG, 'Bold', 640)
    d.text('AI가 직접 고칩니다.', 56, 192, 59, BG, 'Bold', 650)
    d.text('Slow code.\nOptimized by AI.', 58, 295, 39, MINT, 'Regular', 600, 1.15)
    d.text('로우레벨 최적화', 58, 443, 24, LIGHTMUTED, 'SemiBold')
    d.text('Low-level optimization', 268, 445, 24, LIGHTMUTED, 'Regular', 636)
    # Abstract source lines condense into a forward arrow. No invented data.
    for yy, ww in [(121, 142), (145, 109), (169, 157)]:
        d.rect(730, yy, ww, 8, '#47625E', radius=4)
    d.line(798, 207, 798, 226, '#47625E', 3)
    d.mark(742, 269, 128, CORAL)

    # 2 — Explain low-level work through a clearly hypothetical Wine example.
    d.page(False, 'WINE · OPTIMIZATION SCENARIO')
    d.label('Wine 최적화 시나리오', 'Illustrative scenario')
    d.text('같은 결과, 더 적은 작업.', 56, 92, 47, INK, 'Bold', 848)
    d.text('Same result. Less work.', 57, 154, 34, MUTED, 'Regular', 848)
    d.text('Wine: Windows 앱을 Linux에서 실행', 58, 211, 24, MUTED, 'Regular', 840)
    d.text('Runs Windows apps on Linux', 58, 243, 24, MUTED, 'Regular', 840)
    d.rect(56, 300, 403, 161, '#EEEEE7', radius=12)
    d.rect(485, 300, 419, 161, PALE, radius=12)
    d.text('매번 메모리 할당', 79, 321, 29, INK, 'SemiBold', 355)
    d.text('Allocate every time', 80, 359, 26, MUTED, 'Regular', 355)
    d.text('기존 공간 재사용', 511, 321, 29, INK, 'SemiBold', 365)
    d.text('Reuse the memory', 512, 359, 26, MUTED, 'Regular', 365)
    for xx in [82, 188, 294]:
        d.rect(xx, 410, 69, 27, fill='#E0B6A6', radius=5)
        for k in range(3):
            d.line(xx+14+k*16, 417, xx+14+k*16, 430, BG, 2)
    for xx in [152, 258]:
        d.arrow(xx+4, 423.5, xx+26, 423.5, MUTED, 1.8, 5)
    d.rect(514, 410, 76, 27, fill=INK, radius=5)
    for k in range(3):
        d.line(531+k*16, 417, 531+k*16, 430, MINT, 2)
    for xx in [657, 733, 809]:
        d.arrow(604 if xx==657 else xx-44, 423.5, xx-10, 423.5, INK, 1.8, 5)
        d.circle(xx, 412, 22, '#679482')

    # 3 — A continuous agent creates and revises real source changes.
    d.page(False, 'FIND · FIX · RUN')
    d.label('AgenticRocket의 작업', 'How it works')
    d.text('병목을 찾고\n소스를 바꿉니다.', 56, 92, 47, INK, 'Bold', 848, 1.12)
    d.text('Find the bottleneck. Change the code.', 58, 214, 32, MUTED, 'Regular', 848)
    cols = [(56, '01', '병목 찾기', 'Find wasted work'), (357, '02', '코드 수정', 'Rewrite code'), (659, '03', '실행·수정 반복', 'Run. Learn. Retry.')]
    for xx, num, ko, en in cols:
        d.text(num, xx, 294, 40, ORANGE, 'SemiBold', 200)
        d.line(xx, 349, xx+244, 349, RULE, 1.2)
        d.text(ko, xx, 367, 29, INK, 'SemiBold', 255)
        d.text(en, xx, 407, 25, MUTED, 'Regular', 255)
    d.arrow(299, 314, 328, 314, ORANGE, 2.5, 7)
    d.arrow(602, 314, 631, 314, ORANGE, 2.5, 7)
    d.text('한 프로젝트를 끝까지 맡는 AI', 57, 457, 24, INK, 'SemiBold', 430)
    d.text('One agent. One project.', 524, 459, 24, MUTED, 'Regular', 380)

    # 4 — Daytona's role is explicit; three sandboxes are not three hosts.
    d.page(True, 'DAYTONA · EXECUTION & VERIFICATION')
    d.text('Daytona', 56, 33, 28, MINT, 'SemiBold')
    d.text('코드 실행 공간', 220, 40, 20, LIGHTMUTED, 'Regular', 200, essential=False)
    d.text('A workspace for the AI', 370, 41, 20, LIGHTMUTED, 'Regular', 500, essential=False)
    d.text('고친 코드가 정말 빨라졌는지.', 56, 101, 43, BG, 'Bold', 848)
    d.text('Prove the fix makes it faster.', 58, 160, 34, MINT, 'Regular', 848)
    d.text('같은 수정안, 3곳에서 반복 검증', 57, 225, 27, BG, 'SemiBold', 500)
    d.text('One patch. Repeated checks.', 532, 228, 26, LIGHTMUTED, 'Regular', 372)
    for i, xx in enumerate([56, 346, 636]):
        d.rect(xx, 278, 268, 129, PANEL, '#4E6964', radius=10)
        d.text(f'SANDBOX 0{i+1}', xx+22, 294, 15, MINT, 'SemiBold', 220, essential=False)
        d.text('수정 전', xx+22, 326, 27, BG, 'SemiBold', 100)
        d.text('수정 후', xx+160, 326, 27, BG, 'SemiBold', 100)
        d.arrow(xx+111, 341, xx+144, 341, CORAL, 2, 5)
        d.arrow(xx+144, 349, xx+111, 349, CORAL, 2, 5)
        d.text('Before', xx+22, 364, 24, LIGHTMUTED, 'Regular', 100)
        d.text('After', xx+160, 364, 24, LIGHTMUTED, 'Regular', 100)
    d.text('기능 검사와 성능 측정을 함께.', 57, 428, 26, BG, 'SemiBold', 848)
    d.text('Check behavior. Measure speed.', 57, 464, 24, LIGHTMUTED, 'Regular', 848)

    # 5 — Reviewable deliverables, with an editable future screenshot region.
    d.page(False, 'THE PATCH & THE EVIDENCE')
    d.label('개발자에게 남는 것', 'What you take away')
    d.text('수정한 코드와\n검증 근거를 함께.', 56, 104, 42, INK, 'Bold', 425, 1.14)
    d.text('The patch.\nThe proof.', 58, 226, 37, ORANGE, 'Regular', 420)
    d.text('불확실하면 판단 보류', 58, 403, 26, INK, 'SemiBold', 430)
    d.text('Uncertain? No verdict yet.', 58, 439, 25, MUTED, 'Regular', 430)
    d.rect(514, 101, 390, 372, WHITE, RULE, radius=12, name='RESULT_SCREENSHOT_REGION')
    if screenshot:
        d.image(screenshot, 526, 114, 366, 344)
    else:
        d.text('결과 보고서', 540, 120, 23, MUTED, 'SemiBold', 330)
        d.text('Result report', 540, 152, 23, MUTED, 'Regular', 330)
        for yy, ko, en in [(210, '무엇을 바꿨나', 'Code changes'), (300, '기능은 유지됐나', 'Test results'), (390, '얼마나 빨라졌나', 'Measured performance')]:
            d.line(540, yy-16, 878, yy-16, RULE, 1)
            d.text(ko, 542, yy, 29, INK, 'SemiBold', 330)
            d.text(en, 542, yy+39, 25, MUTED, 'Regular', 332)

    # 6 — Close with the actual value proposition, not implementation detail.
    d.page(True, 'LOW-LEVEL PERFORMANCE OPTIMIZATION')
    d.mark(57, 42, 28, CORAL)
    d.text('로우레벨 성능 최적화', 111, 42, 24, LIGHTMUTED, 'SemiBold', 770)
    d.text('Low-level performance optimization', 111, 77, 24, LIGHTMUTED, 'Regular', 770)
    d.text('AgenticRocket', 52, 178, 89, BG, 'Bold', 855)
    d.text('AI가 고치고 실행으로 증명합니다.', 57, 315, 39, BG, 'SemiBold', 848)
    d.text('AI makes the fix. Real runs prove the gain.', 58, 377, 31, MINT, 'Regular', 848)
    d.rect(57, 452, 95, 5, CORAL, radius=2)
    d.save()


def render(screenshot=None):
    out = ROOT/'previews'
    out.mkdir(exist_ok=True)
    pdf = fitz.open(ROOT/'AgenticRocket_3min_KO-EN.pdf')
    thumbs = []
    for i, page in enumerate(pdf):
        pix = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
        path = out/f'slide-{i+1:02d}.png'
        pix.save(path)
        im = Image.open(path).convert('RGB')
        im.thumbnail((640, 360), Image.Resampling.LANCZOS)
        thumbs.append(im)
    contact = Image.new('RGB', (1328, 1160), '#DDE0D9')
    for i, im in enumerate(thumbs):
        contact.paste(im, (16 + (i%2)*656, 16 + (i//2)*384))
    contact.save(out/'overview.png')
    fonts = sorted({f[3] for pg in pdf for f in pg.get_fonts()})
    report = {
        'pages': len(pdf), 'page_points': [W, H], 'aspect_ratio': '16:9',
        'timing_seconds': sum(n['seconds'] for n in NOTES), 'fonts': fonts,
        'fonts_embedded': all(bool(pdf.extract_font(f[0])[3]) for pg in pdf for f in pg.get_fonts() if f[2]=='TrueType'),
        'bilingual_pages': all(any('\uac00' <= c <= '\ud7a3' for c in pg.get_text()) and any(c.isascii() and c.isalpha() for c in pg.get_text()) for pg in pdf),
        'screenshots_used': bool(screenshot),
        'note': 'Wine allocation/reuse is an illustrative optimization scenario, not a measured finding.'
    }
    (ROOT/'verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(json.dumps(report, ensure_ascii=False, indent=2))


def notes():
    parts = ['# AgenticRocket · 3분 발표 대본', '', '화면에는 한국어와 영어가 함께 나옵니다. 발표는 아래 한국어 대본 또는 영어 대본 중 하나를 사용합니다. 두 언어를 연달아 읽으면 3분을 넘습니다.', '', '배정 시간은 총 180초입니다. 문장 사이의 멈춤과 화면을 짚는 시간을 포함한 리허설 기준이며, 실제 발화 속도에 맞춰 조정하세요.', '']
    for lang, name in [('ko', '한국어 발표'), ('en', 'English talk')]:
        parts += [f'## {name}', '']
        for i, n in enumerate(NOTES):
            parts += [f"### {i+1}. {n['title']} · {n['time']} ({n['seconds']}초)", '', n[lang], '']
            if lang == 'ko':
                parts += [f"발표 팁: {n['cue']}", '']
    (ROOT/'speaker-notes.md').write_text('\n'.join(parts), encoding='utf-8')
    (ROOT/'slide-copy.json').write_text(json.dumps(NOTES, ensure_ascii=False, indent=2), encoding='utf-8')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--screenshot', type=Path)
    args = parser.parse_args()
    build(args.screenshot)
    notes()
    render(args.screenshot)
